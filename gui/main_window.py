import tkinter as tk
import platform
import subprocess
from tkinter import ttk, messagebox
from tkinter import PhotoImage
import os 
from ttkthemes import ThemedStyle 
from datetime import datetime
import json

mainFont = "Helvetica"

class MainWindow:
    def __init__(self, root, controller):
        self.root = root
        self.controller = controller
        
        self.controller.set_root(root) 
        self.controller.main_window = self 
        
        self.selected_project_id = None 
        self.last_selected_frame = None
        self.project_image_references = []
        
        self.root.title("ProMan - Project Selector")
        self.root.geometry("1000x600") 

        # --- STYLE DEFINITION ---
        style = ttk.Style(self.root)
        LIST_BG_COLOR = style.lookup('TFrame', 'background') 
        self.bg_color = LIST_BG_COLOR
        
        style.configure("DarkList.TFrame", background=LIST_BG_COLOR)
        style.configure("LeftAnchor.TButton", anchor="w", padding=[1, 1, 1, 1])

        # --- Main Layout Setup ---
        outer_frame = ttk.Frame(self.root, padding="3")
        outer_frame.pack(fill='both', expand=True)

        outer_frame.grid_columnconfigure(0, weight=1) 
        outer_frame.grid_columnconfigure(1, weight=0) 
        outer_frame.grid_rowconfigure(0, weight=1) 

        # --- LEFT SIDE: Project List ---
        self.left_frame = ttk.Frame(outer_frame)
        self.left_frame.grid(row=0, column=0, sticky="nsew") 
        self.left_frame.grid_rowconfigure(1, weight=1) 
        self.left_frame.grid_columnconfigure(0, weight=1) 
        
        ttk.Label(self.left_frame, text="Your Projects (Click to Select):").grid(row=0, column=0, sticky="w", pady=(0, 5))
        
        self.canvas = tk.Canvas(self.left_frame, borderwidth=0, bg=LIST_BG_COLOR, highlightthickness=0)
        self.scrollbar = ttk.Scrollbar(self.left_frame, orient="vertical", command=self.canvas.yview)
        self.project_display_frame = ttk.Frame(self.canvas, style="DarkList.TFrame")
        
        self.canvas.grid(row=1, column=0, sticky="nsew") 
        self.scrollbar.grid(row=1, column=1, sticky="ns")
        
        self.canvas.configure(yscrollcommand=self.scrollbar.set)
        self.canvas_window = self.canvas.create_window((0, 0), window=self.project_display_frame, anchor="nw")
        
        self.canvas.bind('<Configure>', self.on_canvas_resize)
        self.root.bind_all("<MouseWheel>", self._on_mousewheel_propagate)
        
        # --- RIGHT SIDE: Control Buttons ---
        buttons_container = ttk.Frame(outer_frame)
        buttons_container.grid(row=0, column=1, sticky="nw", padx=2, pady=(25, 0)) 

        self.open_project_btn = ttk.Button(buttons_container, text="📂 Open Project", command=self.open_project_clicked, state="disabled", style="LeftAnchor.TButton")
        self.open_project_btn.pack(fill='x', pady=(0, 5), padx=10)

        ttk.Button(buttons_container, text="📂 Edit Attachments", command=self.controller.open_attachment_manager, style="LeftAnchor.TButton").pack(fill='x', pady=(0, 5), padx=10)
        ttk.Button(buttons_container, text="📝 Manage Log Templates", command=self.open_template_manager, style="LeftAnchor.TButton").pack(fill='x', pady=(0, 5), padx=10)
        ttk.Button(buttons_container, text="➕ Create New Project", command=self.controller.open_new_project_dialog, style="LeftAnchor.TButton").pack(fill='x', pady=(0, 5), padx=10) 

        self.edit_project_btn = ttk.Button(buttons_container, text="✏️ Edit Selected Project", command=self.edit_project_clicked, state="disabled", style="LeftAnchor.TButton")
        self.edit_project_btn.pack(fill='x', pady=(0, 5), padx=10)

        self.delete_project_btn = ttk.Button(buttons_container, text="\u232B Delete Selected Project", command=self.delete_project_clicked, state="disabled", style="LeftAnchor.TButton")
        self.delete_project_btn.pack(fill='x', pady=(0, 5), padx=10)

        self.export_docx_btn = ttk.Button(buttons_container, text="📄 Export Project (.docx)", 
                                         command=self.export_project_docx, state="disabled", 
                                         style="LeftAnchor.TButton")
        self.export_docx_btn.pack(fill='x', pady=(0, 5), padx=10)

        self.refresh_project_list() 

    # --- Handlers ---
    def open_template_manager(self):
        from gui.template_manager import TemplateManager
        TemplateManager(self.root, self.controller)

    def open_project_clicked(self):
        if self.selected_project_id:
            self.controller.open_project_detail_window(self.selected_project_id)

    def edit_project_clicked(self):
        if self.selected_project_id:
            self.controller.open_edit_project_dialog(self.selected_project_id)

    def delete_project_clicked(self):
        if self.selected_project_id is not None:
            if messagebox.askyesno("Confirm Deletion", f"Are you sure you want to delete Project ID {self.selected_project_id}?"):
                self.controller.delete_project_flow(self.selected_project_id)

    def on_canvas_resize(self, event):
        self.canvas.itemconfig(self.canvas_window, width=event.width - 5)
        self.project_display_frame.update_idletasks()
        self.canvas.config(scrollregion=self.canvas.bbox("all"))

    def _on_mousewheel_propagate(self, event):
        widget = self.root.winfo_containing(event.x_root, event.y_root)
        while widget:
            if widget == self.left_frame:
                self.on_mousewheel(event)
                return "break"
            widget = widget.master

    def on_mousewheel(self, event):
        delta = event.delta
        scroll_amount = -1 * (delta // abs(delta)) * 4 if abs(delta) >= 120 else -1 * delta * 0.5 
        self.canvas.yview_scroll(int(scroll_amount), "units")

    def select_project(self, project_id: int, frame_widget):
        if self.last_selected_frame:
            try: self.last_selected_frame.config(highlightbackground="#222222")
            except: pass
        
        frame_widget.config(highlightthickness=2, highlightbackground="#FFFFFF")
        self.selected_project_id = project_id
        self.last_selected_frame = frame_widget
        
        self.delete_project_btn.config(state="normal")
        self.open_project_btn.config(state="normal")
        self.edit_project_btn.config(state="normal")
        self.export_docx_btn.config(state="normal")

    def export_project_docx(self):
        if not self.selected_project_id:
            return

        from docx import Document
        from docx.shared import Inches
        from tkinter import filedialog
        import re

        project = self.controller.db_controller.get_project_by_id(self.selected_project_id)
        logs = self.controller.get_all_logs_for_project(self.selected_project_id)
        attachments = self.controller.get_attachments_for_project(self.selected_project_id)
        att_map = {str(a.id): a.file_path for a in attachments}
        
        if not project:
            messagebox.showerror("Error", "Could not find project data.")
            return

        file_path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            initialfile=f"{project.name}_Export.docx",
            filetypes=[("Word Document", "*.docx")]
        )

        if not file_path:
            return

        try:
            doc = Document()

            # --- HEADER SECTION ---
            doc.add_heading(project.name, 0)
            doc.add_paragraph(f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
            
            # --- SEPARATING DESCRIPTION FROM LOGS ---
            description_log = None
            other_logs = []

            for log in logs:
                # We check the timestamp property specifically now
                if log.timestamp == "DESCRIPTION":
                    description_log = log
                else:
                    other_logs.append(log)

            # --- RENDER DESCRIPTION FIRST ---
            if description_log:
                doc.add_heading("Project Description", level=1)
                self._write_log_to_doc(doc, description_log, att_map)
                doc.add_paragraph("") # Space after description
                # If you want a full page break after description, uncomment below:
                # doc.add_page_break() 

            doc.add_paragraph("_" * 30)

            # --- CHRONOLOGICAL LOGS SECTION ---
            # Sort remaining logs by their actual date/time
            sorted_logs = sorted(other_logs, key=lambda x: x.timestamp)

            for log in sorted_logs:
                doc.add_heading(log.timestamp, level=2)
                self._write_log_to_doc(doc, log, att_map)
                doc.add_paragraph("") 

            doc.save(file_path)
            messagebox.showinfo("Success", "Project exported successfully.")

        except Exception as e:
            messagebox.showerror("Export Error", f"An error occurred: {e}")

    def _write_log_to_doc(self, doc, log, att_map):
        """Helper to handle the [ref:x] parsing for any log entry."""
        import re
        from docx.shared import Inches
        
        content_text = ""
        try:
            log_data = json.loads(log.content)
            content_text = log_data.get("text", "").replace('\\n', '\n')
        except:
            content_text = str(log.content)

        if content_text.strip():
            parts = re.split(r'(\[ref:\d+\])', content_text)
            p = doc.add_paragraph()
            
            for part in parts:
                ref_match = re.match(r'\[ref:(\d+)\]', part)
                if ref_match:
                    ref_id = ref_match.group(1)
                    full_path = att_map.get(ref_id)
                    if full_path and os.path.exists(full_path):
                        ext = os.path.splitext(full_path)[1].lower()
                        if ext in ['.png', '.jpg', '.jpeg', '.webp']:
                            doc.add_picture(full_path, width=Inches(4))
                            p = doc.add_paragraph() 
                        else:
                            p.add_run(f" [File: {os.path.basename(full_path)}] ").bold = True
                    else:
                        p.add_run(f" [Missing Ref: {ref_id}] ").italic = True
                else:
                    p.add_run(part)

    def refresh_project_list(self):
        for widget in self.project_display_frame.winfo_children():
            widget.destroy()
        
        self.project_image_references = [] 
        self.selected_project_id = None 
        self.last_selected_frame = None
        
        self.delete_project_btn.config(state="disabled")
        self.open_project_btn.config(state="disabled")
        self.edit_project_btn.config(state="disabled")
        self.export_docx_btn.config(state="disabled")
        
        projects = self.controller.get_all_projects_sorted()

        for p in projects:
            # Border Container
            project_tile_container = tk.Frame(
                self.project_display_frame, 
                bg=self.bg_color,
                highlightthickness=2, 
                highlightbackground="#222222"
            )
            project_tile_container.pack(fill='x', padx=5, pady=5)
            
            inner_content = ttk.Frame(project_tile_container, style="DarkList.TFrame")
            inner_content.pack(fill='both', expand=True, padx=2, pady=2)
            
            # --- RESTORED THUMBNAIL LOGIC ---
            image_label = tk.Label(inner_content, bg=self.bg_color) 
            image_loaded = False 
            if p.thumbnail_path and os.path.exists(p.thumbnail_path):
                try:
                    photo_obj = PhotoImage(file=p.thumbnail_path)
                    image_label.image = photo_obj 
                    image_label.config(image=photo_obj)
                    self.project_image_references.append(photo_obj) 
                    image_loaded = True
                except: pass 

            if not image_loaded:
                image_label.config(text="[No Image]", width=15, height=6, bg="#444", fg="white", font=("Arial", 8)) 
            
            image_label.pack(side='right', padx=10, pady=5)
            
            # --- Text Content (Centered Vertically) ---
            text_frame = ttk.Frame(inner_content, style="DarkList.TFrame")
            text_frame.pack(side='left', fill='both', expand=True, padx=10, pady=5) 
            
            # 1. This is the container that will be centered vertically
            center_group = ttk.Frame(text_frame, style="DarkList.TFrame")
            center_group.pack(expand=True) # expand=True here centers the whole block
            
            priority_text = {3: 'HIGH', 2: 'MEDIUM', 1: 'LOW'}.get(p.priority, 'N/A')
            
            # 2. Pack the Name inside the group (No expand here, so it sits tight)
            name_lbl = ttk.Label(center_group, text=p.name, font=(mainFont, 18, "bold"), 
                                 style="WhiteBold.TLabel")
            name_lbl.pack(side="top", anchor="center")

            # 3. Pack the Details inside the group (No expand here)
            details = f"Priority: {priority_text} | Due: {p.due_date}"
            detail_lbl = ttk.Label(center_group, text=details, font=(mainFont, 12), 
                                   style="White.TLabel")
            detail_lbl.pack(side="top", anchor="center")
            
            # Recursive binding so clicking anything highlights the tile
            def bind_all(widget, pid=p.id, container=project_tile_container):
                widget.bind("<Button-1>", lambda e: self.select_project(pid, container))
                widget.bind("<Double-Button-1>", lambda e: self.open_project_clicked())
                for child in widget.winfo_children():
                    bind_all(child, pid, container)
            bind_all(project_tile_container)

        self.project_display_frame.update_idletasks()
        self.canvas.config(scrollregion=self.canvas.bbox("all"))